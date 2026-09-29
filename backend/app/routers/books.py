import csv
import io
import re
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..auth import get_current_user, require_admin
from ..database import get_db
from ..models import Book, BookWord, Unit, UnitScope, User, Word
from ..schemas import ImportIn, UseIn

router = APIRouter(prefix="/api", tags=["books"])

COLUMN_ALIASES = {
    "text": ["单词", "word", "词语", "英文"],
    "phonetic": ["音标", "phonetic"],
    "pos": ["词性", "pos"],
    "meaning": ["释义", "中文释义", "中文", "意思", "meaning"],
    "example": ["例句", "英文例句", "example"],
    "example_cn": ["例句翻译", "例句译文", "中文例句", "example_cn"],
}


def _get_unit(db: Session, unit_id: int) -> Unit:
    unit = db.get(Unit, unit_id)
    if not unit:
        raise HTTPException(status_code=404, detail="单元不存在，先在书架里选中一个单元")
    return unit


# ---------- 解析 ----------

def _header_map(header: list[str]) -> dict[str, int] | None:
    colmap: dict[str, int] = {}
    for i, h in enumerate(header):
        for key, aliases in COLUMN_ALIASES.items():
            if h.lower() in [a.lower() for a in aliases] and key not in colmap:
                colmap[key] = i
    return colmap if "text" in colmap else None


def _row_item(cells: list[str], colmap: dict[str, int] | None) -> dict:
    if colmap is None:
        # 无表头：按 单词, 音标, 释义, 例句 的顺序（_rest 不含单词本身）
        return {
            "text": cells[0] if cells else "",
            "phonetic": "",
            "meaning": "",
            "example": "",
            "_rest": cells[1:],
        }
    return {k: (cells[i].strip() if i < len(cells) else "") for k, i in colmap.items()}


def _normalize_item(item: dict) -> dict:
    rest = item.pop("_rest", None)
    if rest:
        rest = [c for c in rest if c.strip()]
        if rest and rest[0].startswith("/"):
            if not item.get("phonetic"):
                item["phonetic"] = rest[0]
            rest = rest[1:]
        if rest and not item.get("meaning"):
            item["meaning"] = rest[0]
            rest = rest[1:]
        if rest and not item.get("example"):
            item["example"] = rest[0]
    for k in ("text", "phonetic", "pos", "meaning", "example", "example_cn"):
        item.setdefault(k, "")
    return item


def parse_csv(raw: bytes) -> list[dict]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="文件编码读不出来，请用 Excel 另存为 CSV（UTF-8）后再试")
    rows = [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    if not rows:
        return []
    colmap = _header_map([str(c) for c in rows[0]])
    data_rows = rows[1:] if colmap else rows
    return [_normalize_item(_row_item([str(c).strip() for c in r], colmap)) for r in data_rows]


def parse_xlsx(raw: bytes) -> list[dict]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Excel 文件打不开，确认是 .xlsx 格式")
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [[("" if c is None else str(c).strip()) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    colmap = _header_map(rows[0])
    data_rows = rows[1:] if colmap else rows
    return [_normalize_item(_row_item(r, colmap)) for r in data_rows]


# ---------- 落库 ----------

def _apply_import(db: Session, unit: Unit, items: list[dict]) -> dict:
    existing = {bw.word_id for bw in db.query(BookWord).filter(BookWord.unit_id == unit.id)}
    added, skipped, rows = 0, 0, []

    for it in items:
        text = (it.get("text") or "").strip().lower()
        # 只收像单词的文本，挡掉 aigc:xxx 这类编号或乱码
        if not re.fullmatch(r"[a-z][a-z'\-]{0,29}", text):
            continue
        word = db.query(Word).filter(Word.text == text).first()
        if word is None:
            word = Word(
                text=text,
                phonetic=it.get("phonetic", ""),
                pos=it.get("pos", ""),
                meaning=it.get("meaning") or "（待补）",
                example=it.get("example", ""),
                example_cn=it.get("example_cn", ""),
            )
            db.add(word)
            db.flush()
        else:
            # 词库里已有：补齐空缺字段
            if not word.phonetic and it.get("phonetic"):
                word.phonetic = it["phonetic"]
            if not word.pos and it.get("pos"):
                word.pos = it["pos"]
            if word.meaning == "（待补）" and it.get("meaning"):
                word.meaning = it["meaning"]
            if not word.example and it.get("example"):
                word.example = it["example"]
            if not word.example_cn and it.get("example_cn"):
                word.example_cn = it["example_cn"]

        if word.id in existing:
            skipped += 1
            status = "dup"
        else:
            db.add(BookWord(book_id=unit.book_id, unit_id=unit.id, word_id=word.id, order_no=added))
            existing.add(word.id)
            added += 1
            status = "new"
        rows.append({"word": text, "phonetic": word.phonetic, "meaning": word.meaning, "status": status})

    return {"added": added, "skipped": skipped, "total": added + skipped, "rows": rows[:20]}


# ---------- 接口 ----------

@router.post("/books/{book_id}/use")
def use_book(
    book_id: int,
    body: UseIn | None = None,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    target = user
    uid = body.user_id if body else None
    if uid and uid != user.id:
        if not user.is_admin:
            raise HTTPException(status_code=403, detail="需要管理员权限")
        target = db.get(User, uid)
        if target is None:
            raise HTTPException(status_code=404, detail="账号不存在")
    book = db.get(Book, book_id)
    if not book:
        raise HTTPException(status_code=404, detail="词书不存在")
    target.current_book_id = book.id
    # 指派 = 全选该书的单元（学习 + 练习），清掉旧范围
    db.query(UnitScope).filter(UnitScope.user_id == target.id).delete()
    for u in db.query(Unit).filter(Unit.book_id == book.id):
        db.add(UnitScope(user_id=target.id, unit_id=u.id, study=True, practice=True))
    db.commit()
    return {"ok": True}


@router.post("/lists/import")
def import_words(body: ImportIn, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    unit = _get_unit(db, body.unit_id)
    items = []
    for line in body.text.splitlines():
        cells = [p.strip() for p in line.replace("\t", ",").split(",") if p.strip()]
        if cells:
            items.append(_normalize_item(_row_item(cells, None)))
    result = _apply_import(db, unit, items)
    db.commit()
    return result


@router.post("/lists/import/file")
async def import_file(
    file: UploadFile = File(...),
    unit_id: int = Form(...),
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    unit = _get_unit(db, unit_id)
    raw = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".xlsx"):
        items = parse_xlsx(raw)
    elif name.endswith(".csv") or name.endswith(".txt"):
        items = parse_csv(raw)
    else:
        raise HTTPException(status_code=400, detail="只支持 .xlsx 和 .csv 文件")
    if not items:
        raise HTTPException(status_code=400, detail="文件里没有读到词条，对照模板检查一下格式")
    result = _apply_import(db, unit, items)
    db.commit()
    return result


TEMPLATE_HEADERS = ["单词", "音标", "词性", "释义", "例句", "例句翻译"]
TEMPLATE_EXAMPLES = [
    ["brave", "/breɪv/", "adj.", "勇敢的", "Be brave and try again.", "勇敢一点，再试一次。"],
    ["brilliant", "/ˈbrɪliənt/", "adj.", "聪颖的；杰出的", "She came up with a brilliant idea.", "她想出了一个绝妙的主意。"],
    ["scenery", "/ˈsiːnəri/", "n.", "风景；景色", "The scenery here is beautiful.", "这里的风景很美。"],
]


@router.get("/lists/import/template")
def import_template():
    """导入模板：表头 + 三行示例，第一行表头会被自动识别。"""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "生词导入"
    ws.append(TEMPLATE_HEADERS)
    for row in TEMPLATE_EXAMPLES:
        ws.append(row)
    for i, w in enumerate([12, 14, 8, 18, 36, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = quote("词阶导入模板.xlsx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )
